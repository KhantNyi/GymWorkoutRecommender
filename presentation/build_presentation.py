"""Create editable PPTX and matching vector PDF from one layout definition."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image, ImageDraw
import pymupdf

OUT = Path(__file__).resolve().parent
W, H = 960, 540
GREEN, CREAM, CORAL, LIME = '182B49', 'FFFFFF', '3268C8', '3268C8'
MUTED, BORDER, WHITE, PALE = '53647B', 'DFE6EF', 'FFFFFF', 'F2F6FC'
for name, file in [('Arial', 'segoeui.ttf'), ('Arial-Bold', 'segoeuib.ttf')]:
    pdfmetrics.registerFont(TTFont(name, str(Path('C:/Windows/Fonts') / file)))
prs = Presentation()
prs.slide_width, prs.slide_height = Pt(W), Pt(H)
prs.core_properties.title = 'Gym Workout Recommender'
prs.core_properties.subject = 'Recommendation methods, implementation and validation'
prs.core_properties.author = 'Gym Workout Recommender Project'
pdf = canvas.Canvas(str(OUT / 'Gym_Workout_Recommender.pdf'), pagesize=(W, H))
pdf.setTitle('Gym Workout Recommender')
slide = None
audit = []

def rect(x,y,w,h,fill,stroke=None,round=False):
    round = False
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
    if round: shape.adjustments[0] = .08
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    if stroke: shape.line.color.rgb = RGBColor.from_string(stroke); shape.line.width = Pt(.8)
    else: shape.line.fill.background()
    pdf.setFillColor('#'+fill); pdf.setStrokeColor('#'+(stroke or fill)); pdf.setLineWidth(.8)
    if round: pdf.roundRect(x,H-y-h,w,h,8,stroke=bool(stroke),fill=1)
    else: pdf.rect(x,H-y-h,w,h,stroke=bool(stroke),fill=1)

def text(s,x,y,w,size=19,color=GREEN,bold=False,leading=None):
    font = 'Arial-Bold' if bold else 'Arial'; leading = leading or size*1.3
    lines=[]
    for para in s.split('\n'):
        line=''
        for word in para.split():
            candidate=(line+' '+word).strip()
            if line and pdfmetrics.stringWidth(candidate,font,size)>w:
                lines.append(line); line=word
            else: line=candidate
        lines.append(line)
    for i,line in enumerate(lines):
        yy=y+i*leading
        assert yy+size*1.2 < H, (s,yy)
        assert pdfmetrics.stringWidth(line,font,size)<=w+.5, line
        tb=slide.shapes.add_textbox(Pt(x),Pt(yy),Pt(w+3),Pt(leading+3))
        tf=tb.text_frame; tf.clear(); tf.word_wrap=False
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        p=tf.paragraphs[0]; p.text=line; p.font.name='Segoe UI'; p.font.size=Pt(size)
        p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
        pdf.setFont(font,size); pdf.setFillColor('#'+color)
        pdf.drawString(x,H-yy-size*.92,line)
        audit.append((len(prs.slides),x,yy,w,leading,line))
    return len(lines)*leading

def start(label,title,dark=False,source=''):
    global slide
    if len(prs.slides): pdf.showPage()
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    rect(0,0,W,H,WHITE)
    rect(54,22,26,3,CORAL)
    rect(54,500,852,1,BORDER)
    text(label.upper(),54,32,850,10,CORAL,True)
    if title: text(title,54,66,850,31,GREEN,True)
    text('GYM WORKOUT RECOMMENDER',54,511,300,8,MUTED)
    text(f'{len(prs.slides):02d} / 16',855,511,60,9,MUTED)
    if source: text(source,54,484,835,9,MUTED)

def card(x,y,w,h,title,body,accent=False):
    rect(x,y,w,h,PALE)
    rect(x+18,y+12,22,2,CORAL)
    text(title,x+18,y+23,w-36,18,CORAL if accent else GREEN,True)
    text(body,x+18,y+59,w-36,16,MUTED)

def bullets(items,y=166):
    for item in items:
        rect(56,y+8,6,6,CORAL)
        used=text(item,78,y,810,21)
        y+=used+20

def formula(main,definition):
    rect(54,153,852,100,PALE)
    text(main,75,174,810,25,GREEN)
    text(definition,75,217,810,12,LIME)


def notes(value):
    slide.notes_slide.notes_text_frame.text=value

def band(value,y=425):
    rect(54,y,852,43,PALE)
    text(value,69,y+11,822,16,GREEN)

start('Recommender systems project','')
text('Gym Workout\nRecommender',54,106,650,49,GREEN,True,leading=59)
text('From lecture techniques to a practical workout plan',57,247,790,23,MUTED)
rect(57,310,842,1,BORDER)
for x,chapter,label in [(57,'CHAPTER 4','Content-based'),(346,'CHAPTER 8','Context-aware'),(635,'CHAPTER 10','Weighted hybrid')]:
    text(chapter,x,335,245,11,CORAL,True)
    text(label,x,361,245,22,GREEN,True)
text('Streamlit application • Local exercise catalogue • Explainable ranking',57,443,835,15,MUTED)
notes('Presentation grounded in the course materials in C:/Users/ADMIN/Downloads/Recommender and the current app.py, engine.py and storage.py. The application implements exactly three selectable methods: Content-based, Context-aware and Hybrid. All exact gym scoring weights are project defaults.')

start('01 / Project purpose','Match the workout to the user.',source='Project proposal, slide 2; app.py; engine.py')
card(54,151,274,230,'Training goals','General Fitness\nWeight Loss\nMuscle Gain\nStrength')
card(343,151,274,230,'Practical inputs','Experience\nEquipment and muscle focus\nAvailable minutes\nLocation and current energy')
card(632,151,274,230,'Useful output','Ranked exercises\nA time-budgeted plan\nReasons for each selection\nFeedback for later sessions')
band('Objective: personalized exercise selection with clear, inspectable ranking reasons.')

start('02 / App scope','Three methods. One clear workout builder.',source='app.py — Recommendation method; engine.py — METHODS')
for x,a,b in [(54,'Content-based','Matches exercise attributes to your goal and liked exercises.'),(343,'Context-aware','Uses goal/experience rules and personal completion history.'),(632,'Hybrid','Combines content, rating evidence, knowledge and context.')]:
    card(x,159,274,211,a,b)
text('Goal and experience are separate choices.',54,397,840,21,GREEN,True)
band('General Fitness / Weight Loss / Muscle Gain / Strength  •  Beginner / Intermediate / Advanced')
notes('Popularity and Knowledge remain internal components of Hybrid, not separate methods. SVD, User CF, Item CF and case retrieval were removed. Advanced is the UI label for the dataset expert level.')

start('03 / Data','One catalogue, separate sources of evidence.',source='docs/DATA_SOURCES.md; reports/data_quality.json; storage.py')
for x,num,label in [(54,'2,872','unique exercise entries'),(343,'580','instruction matches'),(632,'24','exercises with local photos')]:
    text(num,x,152,270,46,CORAL,True); text(label,x,214,270,17,MUTED)
card(54,269,411,176,'Exercise catalogue','Kaggle exercise descriptions, muscles, equipment, levels and ratings.\nDetailed instructions are matched from Free Exercise DB.')
card(483,269,423,176,'Feedback and demonstrations','SQLite stores exercise ratings and completion.\n48 local photos illustrate 24 matched exercises. Missing images are clearly marked.')
notes('Sources: Kaggle gym-exercise-data (niharika41298); Kaggle gym-members-exercise-dataset (valakhorasani); github.com/yuhonas/free-exercise-db. The catalogue was deduplicated from 2,918 rows. Member-session data remains archived and is not used in the simplified app. Demonstration files are mapped by exact normalized exercise name, with hashes and original URLs in data/images/manifest.json.')

start('04 / Lecture alignment','Where each part of the app comes from.',source='Supplied Chapters 3, 4, 7, 8 and 10; engine.py; docs/CHAPTER_MAPPING.md')
rows=[['Lecture topic','Application use'],['Ch. 4 — Cosine similarity','Content-based; TF-IDF is our representation choice'],['Ch. 7 — Knowledge-based','Shared hard constraints and goal/experience rules'],['Ch. 8 — Context-aware','Workout context; completion-based score adaptation'],['Ch. 10 — Weighted hybrid','Weighted sum of component scores'],['Internal rating evidence','Ch. 3 popularity prior; no separate selectable mode']]
for i,row in enumerate(rows):
    for x,w,val in [(54,315,row[0]),(369,537,row[1])]:
        rect(x,145+i*49,w,49,PALE if i%2==0 else WHITE)
        text(val,x+13,159+i*49,w-26,15,GREEN,i==0)
notes('Chapter 4 code cells 33–46 use CountVectorizer and cosine_similarity. Chapter 7 includes hard-constraint filtering and case similarity. Chapter 8 includes availability filtering and relevance-based scoring. Chapter 10 cells 35 and 55 provide weighted hybrid functions. The current app removes collaborative methods and uses one fixed weighted hybrid formula.')

start('05 / Content-based','Same similarity principle, exercise features.',source='Chapter 4: cosine_similarity examples and movie-feature code; engine.py: content scoring')
formula('Similarity(q, x) = (q · x) / (||q|| × ||x||)','q = goal/focus vector; x = exercise vector')
card(54,275,411,182,'In the lecture','Combine movie attributes, build count vectors and compare them with cosine similarity.\nRank movies similar to a liked movie.')
card(483,275,423,182,'In our app','TF-IDF encodes exercise attributes. Match goal and muscle focus.\nWith ratings ≥ 4: 60% profile similarity + 40% liked-exercise similarity.')
notes('Chapter 4 uses CountVectorizer over keywords, cast, genres and director. Our adaptation uses TfidfVectorizer with word unigrams/bigrams over name, description, category, muscle, level and equipment. The 60/40 blend is a project choice.')

start('06 / Context','Adapt the context concept to a workout.',source='Chapter 8: time/location-aware examples; engine.py: constraints, Context and make_plan')
card(54,151,411,235,'In the lecture','Filter items by time or location availability.\n\nThe time-aware food example ranks by preference score × time relevance.\n\nContext changes what is suitable.')
card(483,151,423,235,'In our app','Location restricts equipment; low energy limits difficulty.\n\nThe Context score combines knowledge rules with personal completion history.\n\nTime budget controls plan allocation.')
band('Conceptual adaptation: the app does not use the lecture’s time-of-day scoring formula.')
notes('Chapter 8 cells 23–65 determine meal context, filter availability and multiply Preference_Score by Time_Relevance. The app uses Gym/Home/Outdoors, energy and a session budget; it does not use GPS distance or meal-time categories. These constraints apply to all ranking methods.')

start('07 / Knowledge and constraints','Apply feasibility rules across every method.',source='Chapter 7: constraint-based recommendation; engine.py: recommend and make_plan')
for i,(a,b) in enumerate([('Equipment','Required equipment must be available; bodyweight is always included.'),('Experience','Reject exercises above the experience limit; low energy caps it at beginner.'),('Muscle preferences','Respect focus and excluded primary muscles.'),('Diversity','Subtract 0.10 per already-selected exercise with the same primary muscle.'),('Time allocation','Reserve 5 minutes; allocate 5- or 7-minute exercise blocks as applicable.')]):
    y=148+i*61
    text(f'{i+1:02d}',54,y,45,20,CORAL,True)
    text(a,112,y,200,18,GREEN,True); text(b,318,y,580,16,MUTED)
notes('The time block is min(5 for low energy, otherwise 7, budget minus 5), allowing a 10-minute session. The ranked list is greedily diversified, then make_plan selects the number that fits. These are project planning rules, not lecture-prescribed exercise durations.')

start('08 / Weighted hybrid','Chapter 10: add weighted score contributions.',source='Chapter 10: weighted_hybrid and hybrid_recommender examples, code cells 35–57')
formula('Hybrid score = Σ weight × component score','Each component contributes to a shared score; example weights sum to 1.')
card(54,275,411,182,'Two-component lecture example','Hybrid = 0.60 CF + 0.40 Content\n\nFor CF = 4.8 and Content = 4.2:\nHybrid = 2.88 + 1.68 = 4.56.')
card(483,275,423,182,'Generalized lecture example','Hybrid = 0.50 CF + 0.30 Content + 0.20 Context\n\nThe lecture also accepts a dictionary of score columns and weights.')
notes('Lecture examples use movie scores. Cells 20–27 specify 0.6/0.4 and compute the weighted sum. Cells 43–57 add Context and a general dictionary-based function. These examples establish the mechanism, not optimal gym weights.')

start('09 / Our weighted hybrid','The lecture mechanism, adapted to gym signals.',source='engine.py: Hybrid score; Chapter 10: generalized weighted sum')
parts=[('Content','Goal/focus and liked items',.40,'3268C8'),('Popularity','Source prior and local ratings',.20,'5A87D5'),('Knowledge','Goal/category and experience',.25,'86A6E0'),('Context','Knowledge plus completion',.15,'B6CBEF')]
xx=54
for label,desc,weight,color in parts:
    rect(xx,151,852*weight,41,color)
    text(f'{weight:.0%}',xx+13,160,852*weight-20,18,WHITE if weight>=.2 else GREEN,True)
    xx+=852*weight
for i,(label,desc,weight,color) in enumerate(parts):
    yy=223+i*44
    text(label,54,yy,200,18,GREEN,True); text(desc,260,yy,640,18,MUTED)
band('B = 0.40C + 0.20P + 0.25K + 0.15X   •   Weights are project defaults.')
notes('C is content cosine. P=(5*source_prior + rating_count*local_mean)/(5+rating_count), using normalized rating scales. K=0.7*category_goal_match+0.3*experience_match. X=0.6*K+0.4*personal_mean_completion, default 0.5 completion. Context reuses Knowledge, so the components are not independent.')

start('10 / Worked example','Inspect how each signal contributes.',source='Illustrative inputs using the actual engine.py formula; these are not measured recommendations.')
rows=[['Signal','Input','Weight','Contribution'],['Content C','0.60','0.40','0.24'],['Popularity P','0.70','0.20','0.14'],['Knowledge K','1.00','0.25','0.25'],['Context X','0.80','0.15','0.12']]
for i,row in enumerate(rows):
    for j,value in enumerate(row):
        x=54+j*213
        rect(x,150+i*45,213,45,PALE if i%2==0 else WHITE)
        text(value,x+16,162+i*45,180,17,GREEN,i==0)
text('Base hybrid score',54,405,650,24,GREEN,True)
text('0.75',737,389,169,44,CORAL,True)
notes('Example: 0.40*0.60+0.20*0.70+0.25*1.00+0.15*0.80=0.75. K=1 and default completion H=.5 give X=.8. These scores are not probabilities and the coefficients are not trained weights.')

start('11 / Feedback','Feedback changes scores, not the selected method.',source='engine.py — content, rating prior, completion history and fixed Hybrid formula')
card(54,151,411,245,'With no exercise history','Content uses the goal/focus query.\nPopularity uses a source-rating prior.\nCompletion defaults to 0.5.\n\nAll three methods can produce a workout from the profile alone.')
card(483,151,423,245,'After logging activity','Ratings of 4 or 5 inform liked-exercise similarity.\nLocal ratings update rating evidence.\nCompletion updates Context.\n\nHybrid keeps the same 40/20/25/15 weights.')
band('Weighted hybrid from Chapter 10; the lecture’s 5/20-rating switching rule is not used.')
notes('No collaborative blend or CF fallback remains in the simplified app. Feedback updates the component values. The lecture switching example is discussed only to distinguish it from the implemented fixed weighted sum.')

start('12 / Method comparison','Different ranking signals, shared constraints.',source='engine.py — implementation comparison, not a measured quality ranking')
rows=[['Question','Content-based','Context-aware','Hybrid'],['Primary signal','Attribute similarity','Rules + completion','Weighted scores'],['Works without ratings?','Yes','Yes; default history','Yes; base blend'],['Uses liked exercises?','Yes, when available','No rating-based term','Via content component'],['Internal components','Cosine similarity','Knowledge + history','Content + P + K + X'],['Equipment / time rules','Shared','Shared','Shared']]
for i,row in enumerate(rows):
    widths=[234,206,206,206]; xx=54
    for j,value in enumerate(row):
        rect(xx,150+i*48,widths[j],48,PALE if i%2==0 else WHITE)
        text(value,xx+11,164+i*48,widths[j]-22,14,GREEN,i==0)
        xx+=widths[j]

start('13 / Live demonstration','Show the selected method changing the workout.',source='app.py — Build workout, Exercise library, Activity')
for i,(a,b) in enumerate([('Save a profile','Set goal, experience, equipment and time; click Build my workout.'),('Select a method','Use the three visible Recommendation method choices in the builder.'),('Compare the three','Switch between Content-based, Context-aware and Hybrid.'),('Inspect the evidence','Read How to perform, photos and Why this exercise? details.'),('Record feedback','Use Log this exercise, then save feedback in Activity.')]):
    yy=147+i*62
    text(f'{i+1:02d}',54,yy,45,23,CORAL,True)
    text(a,112,yy,220,18,GREEN,True); text(b,352,yy,552,16,MUTED)
notes('Pending filter edits are explicitly marked; only Build my workout applies them. Method switching uses the applied profile. How the ranking works reveals component scores. The library supports search and a photos-only filter. Activity logging prefills the selected exercise. CSV and JSON downloads remain available.')

start('14 / Validation','What has been checked—and what remains open.',source='Current Streamlit suite: 28 tests passed; recorded catalogue checks: reports/evaluation.json')
for x,num,label in [(54,'28','application tests passed'),(343,'36','catalogue scenarios checked'),(632,'0','equipment / budget violations')]:
    text(num,x,146,270,46,CORAL,True); text(label,x,211,270,15,MUTED)
card(54,265,411,185,'Supported by checks','Three scoring methods, time budgets, migration and image integrity.\n\nUI checks cover method changes, draft filters, navigation and logging.')
card(483,265,423,185,'Still to evaluate','Recall@10 and NDCG@10 require eligible genuine held-out histories.\n\nNo measured best-method result or fitness-outcome validation.')
notes('The current app suite passed 28 tests, including media integrity and migration. Desktop/mobile browser checks confirmed navigation, method selection and local image loading. The rerun 36 scenarios had zero budget/equipment violations and 0.45% catalogue coverage for the narrow tested set. One upstream deprecation warning remains.')

start('15 / Conclusion','Lecture foundations. A gym-specific application.')
for yy,a,b in [(151,'Content-based','Applies Chapter 4 similarity to exercise attributes.'),(221,'Context-aware','Adapts Chapter 8 context to workout feasibility and history.'),(291,'Weighted hybrid','Uses Chapter 10 score blending with project-specific weights.')]:
    text(a,54,yy,245,22,CORAL,True); text(b,325,yy,575,20,GREEN)
rect(54,390,852,74,PALE)
text('Next: collect genuine feedback, then evaluate and tune the weights.',73,411,815,21,GREEN,True)
notes('Key distinction for the presentation: the mathematical weighted-hybrid mechanism corresponds to Chapter 10; the component definitions, weights, feedback definitions, diversity penalty and workout allocation are project adaptations. The app is an educational prototype.')

prs.save(OUT/'Gym_Workout_Recommender.pptx')
pdf.save()
doc=pymupdf.open(OUT/'Gym_Workout_Recommender.pdf')
thumbs=[]
for i,page in enumerate(doc):
    pix=page.get_pixmap(matrix=pymupdf.Matrix(1,1),alpha=False)
    pix.save(OUT/f'slide_{i+1:02d}.png')
    im=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
    im.thumbnail((480,270)); thumbs.append(im)
contact=Image.new('RGB',(1440,6*294),'#E5E5E5'); d=ImageDraw.Draw(contact)
for i,im in enumerate(thumbs):
    x=(i%3)*480; y=(i//3)*294
    contact.paste(im,(x,y)); d.text((x+9,y+274),f'Slide {i+1:02d}',fill='#10131C')
contact.save(OUT/'slide_overview.jpg',quality=92)
assert len(doc)==len(prs.slides)==16
assert all(page.get_text().strip() for page in doc)
(OUT/'README.md').write_text('''# Gym Workout Recommender presentation

- `Gym_Workout_Recommender.pptx`: 16 editable widescreen slides.
- `Gym_Workout_Recommender.pdf`: matching vector PDF.
- `slide_overview.jpg`: overview of all slides.
- `build_presentation.py`: reproducible layout and content source.

Clean white presentation with blue accents, navy typography, pale panels and generous spacing. The 16 slides map the current Streamlit implementation to Chapters 4, 7, 8 and 10, with internal rating and knowledge components acknowledged. It explicitly distinguishes the lecture examples from project-specific representations, context rules, weights and feedback behavior. The app has exactly three methods. General Fitness is a goal; Beginner is experience only. The fixed hybrid contains rating-prior and knowledge components but no collaborative blend. All worked calculations are illustrative. Lecture sources are identified by chapter and code-cell ranges in speaker notes. The deck describes only the Streamlit application.

Validation: the app suite has 28 passing tests, with desktop/mobile browser checks and local image loading verified. The 36 catalogue scenarios were rerun successfully.

The PowerPoint uses native editable text and shapes. The PDF and previews are exported from the same layout source. Run `.venv/Scripts/python.exe presentation/build_presentation.py` from the project root to rebuild (requires python-pptx, reportlab, pymupdf and Pillow).
''',encoding='utf-8')
print(f'Created {len(prs.slides)} slides; {len(audit)} text lines checked; PPTX and PDF saved in {OUT}')
